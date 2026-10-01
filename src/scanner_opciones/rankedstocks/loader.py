"""Lectura del ranking que el usuario descarga de RankedStocks (.xlsx). Funciones puras, sin red.

La app NO consulta rankedstocks.com (sus términos prohíben el scraping sin permiso escrito): solo lee
el fichero que el usuario elige. Las columnas se infieren del propio fichero para no romperse si el
formato del export cambia un poco.
"""
from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Optional

from scanner_opciones.domain.errors import WatchlistError

MAX_CHOICES = 12                                    # una columna con ≤ 12 valores distintos se filtra por lista
SYMBOL_HEADERS = {"simbolo", "symbol", "symbols", "ticker", "tickers"}
_FLAG = re.compile("[\U0001F1E6-\U0001F1FF]")       # banderas: pares de «regional indicator»
_ARROWS = re.compile(r"\s*[↑↓▲▼⬆⬇]+\s*$")           # «RS ↓»: la flecha es del orden en la web de origen
_NUMBER = re.compile(r"^[+-]?\d+(\.\d+)?$")
_SUFFIX = {"K": 1e3, "M": 1e6, "B": 1e9, "T": 1e12}

TEXT, CHOICE, NUMBER = "text", "choice", "number"


@dataclass(frozen=True)
class Column:
    name: str
    kind: str                                       # text | choice | number
    choices: tuple[str, ...] = ()                   # valores distintos (solo kind == choice)


@dataclass(frozen=True)
class Cell:
    text: str                                       # tal como viene en el fichero
    value: Optional[float] = None                   # valor numérico si la columna es numérica


@dataclass(frozen=True)
class Row:
    ticker: str
    cells: tuple[Cell, ...]


@dataclass(frozen=True)
class RankedTable:
    source: str                                     # nombre del fichero
    columns: tuple[Column, ...]
    rows: tuple[Row, ...]
    symbol_index: int


def parse_number(text: str) -> Optional[float]:
    """«$4.4B» -> 4.4e9, «$71.37» -> 71.37, «95.6» -> 95.6, «1,234» -> 1234, «12%» -> 12. None si no es un número."""
    t = (text or "").strip().replace("$", "").replace(",", "").replace("%", "").replace(" ", "")
    if not t:
        return None
    mult = 1.0
    if t[-1].upper() in _SUFFIX:
        mult = _SUFFIX[t[-1].upper()]
        t = t[:-1]
    if not _NUMBER.match(t):
        return None
    return float(t) * mult


def clean_header(name: object, index: int) -> str:
    text = _ARROWS.sub("", str(name).strip()) if name is not None else ""
    return text or f"Columna {index + 1}"


def clean_ticker(text: str) -> str:
    """«🇺🇸DK» -> «DK»."""
    return _FLAG.sub("", text).strip().upper()


def _plain(text: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", text.lower()) if unicodedata.category(c) != "Mn")


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def load_table(path: str | Path) -> RankedTable:
    p = Path(path)
    if not p.is_file():
        raise WatchlistError(f"No existe el archivo: {p}")
    if p.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise WatchlistError(f"Formato no soportado: {p.suffix or '(sin extensión)'} (el ranking debe ser un .xlsx)")
    try:
        from openpyxl import load_workbook

        wb = load_workbook(p, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl lanza tipos variados con archivos corruptos
        raise WatchlistError(f"No se pudo abrir el Excel {p.name}: {exc}") from exc
    try:
        raw = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    finally:
        wb.close()
    return build_table(p.name, raw)


def build_table(source: str, raw: list[list[object]]) -> RankedTable:
    """`raw`: la primera fila es la cabecera. Separado de `load_table` para poder probarlo sin ficheros."""
    raw = [r for r in raw if any(_cell_text(c) for c in r)]
    if len(raw) < 2:
        raise WatchlistError("El archivo no contiene filas de datos")
    width = max(len(r) for r in raw)
    names = [clean_header(raw[0][i] if i < len(raw[0]) else None, i) for i in range(width)]
    symbol = next((i for i, n in enumerate(names) if _plain(n) in SYMBOL_HEADERS), None)
    if symbol is None:
        raise WatchlistError("No se encuentra la columna «Símbolo» en el archivo")
    data = [[_cell_text(r[i]) if i < len(r) else "" for i in range(width)] for r in raw[1:]]
    data = [r for r in data if clean_ticker(r[symbol])]

    columns: list[Column] = []
    numeric: list[bool] = []
    for i, name in enumerate(names):
        values = [r[i] for r in data if r[i]]
        is_number = i != symbol and bool(values) and all(parse_number(v) is not None for v in values)
        numeric.append(is_number)
        distinct = sorted(set(values))
        if is_number:
            columns.append(Column(name, NUMBER))
        elif i != symbol and 0 < len(distinct) <= min(MAX_CHOICES, max(1, len(data) // 2)):
            columns.append(Column(name, CHOICE, tuple(distinct)))
        else:
            columns.append(Column(name, TEXT))
    rows = tuple(
        Row(
            clean_ticker(r[symbol]),
            tuple(Cell(r[i], parse_number(r[i]) if numeric[i] else None) for i in range(width)),
        )
        for r in data
    )
    if not rows:
        raise WatchlistError("El archivo no contiene filas con símbolo")
    return RankedTable(source, tuple(columns), rows, symbol)
