"""Carga de watchlists desde archivos .txt / .csv / .xlsx."""
from __future__ import annotations

from pathlib import Path

from scanner_opciones.domain.errors import WatchlistError
from scanner_opciones.watchlist.parser import HEADER_NAMES, ParseResult, parse_text, parse_tokens

TEXT_SUFFIXES = {".txt", ".csv"}
EXCEL_SUFFIXES = {".xlsx", ".xlsm"}


def load_watchlist_file(path: str | Path) -> ParseResult:
    p = Path(path)
    if not p.is_file():
        raise WatchlistError(f"No existe el archivo: {p}")
    suffix = p.suffix.lower()
    if suffix in TEXT_SUFFIXES:
        return _load_text(p)
    if suffix in EXCEL_SUFFIXES:
        return _load_excel(p)
    raise WatchlistError(f"Formato no soportado: {suffix or '(sin extensión)'} (usa .txt, .csv o .xlsx)")


def _load_text(p: Path) -> ParseResult:
    try:
        text = p.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = p.read_text(encoding="latin-1")
        except OSError as exc:
            raise WatchlistError(f"No se pudo leer {p}: {exc}") from exc
    except OSError as exc:
        raise WatchlistError(f"No se pudo leer {p}: {exc}") from exc
    return parse_text(text)


def _load_excel(p: Path) -> ParseResult:
    """Lee la primera columna de la primera hoja; una primera fila de cabecera se ignora."""
    try:
        from openpyxl import load_workbook

        wb = load_workbook(p, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl lanza tipos variados con archivos corruptos
        raise WatchlistError(f"No se pudo abrir el Excel {p}: {exc}") from exc
    try:
        ws = wb.worksheets[0]
        tokens: list[str] = []
        for i, row in enumerate(ws.iter_rows(min_col=1, max_col=1, values_only=True)):
            value = row[0] if row else None
            if value is None:
                continue
            text = str(value).strip()
            if i == 0 and text.upper() in HEADER_NAMES:
                continue
            tokens.append(text)
    finally:
        wb.close()
    return parse_tokens(tokens)
