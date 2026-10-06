"""Universo de acciones: une los ficheros que el usuario descarga de RankedStocks y HelloStocks.

Cada fuente tiene sus propias columnas: RankedStocks es un fichero con una hoja (fuente «RankedStocks») y
HelloStocks un libro con una pestaña por estrategia (la fuente es el nombre de la pestaña). La app no consulta
ninguna web: solo lee los ficheros elegidos. Funciones puras, sin red.
"""
from __future__ import annotations

from dataclasses import dataclass, replace
from pathlib import Path

from scanner_opciones.domain.errors import WatchlistError
from scanner_opciones.rankedstocks.loader import (
    SYMBOL_HEADERS, TEXT, RankedTable, _cell_text, _plain, build_table, clean_header, clean_ticker,
)

RANKED = "RankedStocks"
SOURCE_COLUMN = "Fuente"
ALL = ""                                           # valor de `src` en la URL para la vista «Todas»
# La pestaña «Value Investing-Warren Buffett» de HelloStocks viene sin fila de cabecera; estas columnas se
# deducen comparando sus valores con los de las otras pestañas (mismo ticker, mismos números).
HEADERLESS_COLUMNS = (
    "Ticker", "Company", "Sector", "Criteria", "Revenue Growth (5Y)", "ROE", "Debt to Equity",
    "Free Cash Flow (TTM)", "PE Ratio", "PB Ratio", "Dividend Yield",
)
_NAME_HEADERS = {"company", "empresa"}


@dataclass(frozen=True)
class Source:
    name: str                                      # «RankedStocks» o el nombre de la pestaña de HelloStocks
    file: str                                      # fichero del que sale
    table: RankedTable                             # con la columna «Fuente» justo detrás del símbolo


def _is_ranked(header: list[object]) -> bool:
    """Un export de RankedStocks lleva la columna «RS» (el nombre del fichero no es fiable)."""
    return any(_plain(clean_header(h, i)) == "rs" for i, h in enumerate(header))


def _with_source(name: str, raw: list[list[object]], symbol: int) -> list[list[object]]:
    """Inserta la columna «Fuente» (constante) detrás del símbolo."""
    out = []
    for n, row in enumerate(raw):
        row = list(row) + [None] * (symbol + 1 - len(row))
        out.append(row[: symbol + 1] + [SOURCE_COLUMN if n == 0 else name] + row[symbol + 1:])
    return out


def sheet_to_source(file: str, sheet: str, rows: list[list[object]], ranked_name: str = RANKED) -> Source:
    rows = [r for r in rows if any(_cell_text(c) for c in r)]
    if not rows:
        raise WatchlistError(f"La hoja «{sheet}» está vacía")
    first = [clean_header(c, i) for i, c in enumerate(rows[0])]
    if not any(_plain(h) in SYMBOL_HEADERS for h in first):      # sin cabecera: se usa la deducida si cuadra
        if len(first) != len(HEADERLESS_COLUMNS):
            raise WatchlistError(f"La hoja «{sheet}» no tiene cabecera con «Ticker»")
        rows = [list(HEADERLESS_COLUMNS), *rows]
    symbol = next(i for i, h in enumerate(rows[0]) if _plain(clean_header(h, i)) in SYMBOL_HEADERS)
    name = ranked_name if _is_ranked(rows[0]) else sheet.strip()
    return Source(name, file, build_table(name, _with_source(name, rows, symbol)))


def load_sources(path: str | Path, file: str | None = None) -> list[Source]:
    """Todas las hojas del .xlsx como fuentes (las hojas sin datos útiles se ignoran, salvo que ninguna sirva)."""
    p = Path(path)
    if not p.is_file():
        raise WatchlistError(f"No existe el archivo: {p}")
    if p.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise WatchlistError(f"Formato no soportado: {p.suffix or '(sin extensión)'} (el universo debe ser un .xlsx)")
    try:
        from openpyxl import load_workbook

        wb = load_workbook(p, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl lanza tipos variados con archivos corruptos
        raise WatchlistError(f"No se pudo abrir el Excel {file or p.name}: {exc}") from exc
    try:
        sheets = [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]
    finally:
        wb.close()
    sources: list[Source] = []
    errors: list[str] = []
    for title, rows in sheets:
        try:
            sources.append(sheet_to_source(file or p.name, title, rows))
        except WatchlistError as exc:
            errors.append(str(exc))
    if not sources:
        raise WatchlistError(errors[0] if errors else "El archivo no contiene hojas")
    return sources


def merge(sources: list[Source]) -> RankedTable:
    """Vista «Todas»: una fila por ticker con su empresa, sector y las fuentes en las que aparece."""
    info: dict[str, dict] = {}
    for src in sources:
        names = [_plain(c.name) for c in src.table.columns]
        company = next((i for i, n in enumerate(names) if n in _NAME_HEADERS), None)
        sector = next((i for i, n in enumerate(names) if n == "sector"), None)
        for row in src.table.rows:
            entry = info.setdefault(row.ticker, {"company": "", "sector": "", "sources": []})
            if company is not None and not entry["company"]:
                entry["company"] = row.cells[company].text
            if sector is not None and not entry["sector"]:
                entry["sector"] = row.cells[sector].text
            if src.name not in entry["sources"]:
                entry["sources"].append(src.name)
    raw: list[list[object]] = [["Ticker", "Empresa", "Sector", "Fuentes", "Nº fuentes"]]
    for ticker in sorted(info):
        e = info[ticker]
        raw.append([ticker, e["company"], e["sector"], " · ".join(e["sources"]), len(e["sources"])])
    if len(raw) < 2:
        raise WatchlistError("El universo está vacío")
    table = build_table("Todas", raw)
    # «Fuentes» siempre se filtra por texto (con pocos valores distintos se inferiría una lista de combinaciones)
    columns = tuple(replace(c, kind=TEXT, choices=()) if c.name == "Fuentes" else c for c in table.columns)
    return replace(table, columns=columns)
